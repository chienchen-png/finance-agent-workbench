"""Phase 17 P11 集成验证：公式兜底重算（v6.2 默认禁用）。

构造"第三方生成、无缓存值"的 xlsx（公式写进去但 data_only 读为 None），
验证：
  1. **默认（FINMOD_FORMULA_RECALC 未设）**：跳过重算，公式列保持 None
     （导入秒级完成，防 formulas 全簿重算卡死——2026-08-21 v6.2 修复）
  2. **显式启用（FINMOD_FORMULA_RECALC=1）**：formulas 重算回填公式列
  3. 正常文件（有缓存值）不受影响
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import storage.db as sdb  # noqa: E402
_tmp = tempfile.mkdtemp(prefix="ff_")
sdb.DB_PATH = os.path.join(_tmp, "test.db")

from app import create_app  # noqa: E402
from storage.project_store import ProjectStore  # noqa: E402
from tools.project_data_importer import ProjectDataImporter  # noqa: E402

FAIL = []


def check(label: str, ok: bool, detail: str = "") -> None:
    print(("PASS" if ok else "FAIL"), label, detail)
    if not ok:
        FAIL.append(label)


def main() -> int:
    app = create_app()
    with app.app_context():
        from storage.db import get_db
        db = get_db()
        work = os.path.join(_tmp, "work")
        os.makedirs(work, exist_ok=True)

        # ── 1. 构造无缓存值 xlsx：公式写入但无缓存值 ──
        import openpyxl
        p = os.path.join(work, "公式测试.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "回款表"
        ws.append(["合同编号", "金额A", "金额B", "合计"])
        ws.append(["HT-001", 100, 200, "=B2+C2"])
        ws.append(["HT-002", 300, 400, "=B3+C3"])
        ws.append(["HT-003", 500, 600, "=B4+C4"])
        wb.save(p)
        wb.close()

        # 确认：data_only 读为 None（无缓存值）
        wb2 = openpyxl.load_workbook(p, data_only=True, read_only=True)
        d = wb2["回款表"]["D2"].value
        wb2.close()
        check("预置：公式无缓存值（data_only=None）", d is None, f"值={d!r}")

        # ── 2. 默认禁用：导入 → 跳过重算（公式列保持 None，秒级完成） ──
        proj = ProjectStore(db).create("公式验证", work)
        imp = ProjectDataImporter(db)
        r = imp.import_workbook(proj["id"], "公式测试.xlsx", work)
        check("导入成功（默认禁用，秒级）", r.get("file_id") is not None,
              str(r.get("sheets")))

        # ── 3. 数据平面验证：公式列保持 None（v6.2 默认禁用） ──
        from storage.project_data_store import ProjectDataStore
        store = ProjectDataStore(db)
        sheets = store.list_sheets(r["file_id"])
        ss = next(s for s in sheets if s["sheet_name"] == "回款表")
        headers = [h["name"] for h in ss["headers"]]
        print("  列名:", headers)
        rows = db.execute(
            f'SELECT * FROM "{ss["table_name"]}" ORDER BY id').fetchall()
        check("3 行数据", len(rows) == 3, f"实际 {len(rows)}")
        if len(rows) == 3:
            # 列名可能是 合计 或 合计_2 等
            total_col = next((h for h in headers if "合计" in h), headers[-1])
            vals = [dict(x)[total_col] for x in rows]
            check("默认禁用：合计列保持 None（跳过重算）",
                  all(v is None for v in vals), f"合计={vals}（列={total_col}）")

        # ── 3b. 显式启用 FINMOD_FORMULA_RECALC=1：重算回填 ──
        # 用独立新项目（同 hash 幂等复用会跳过重算，新项目触发全量导入）
        os.environ["FINMOD_FORMULA_RECALC"] = "1"
        try:
            proj3 = ProjectStore(db).create("公式验证-启用", work)
            r3 = imp.import_workbook(proj3["id"], "公式测试.xlsx", work)
            sheets3 = store.list_sheets(r3["file_id"])
            ss3 = next(s for s in sheets3 if s["sheet_name"] == "回款表")
            h3 = [h["name"] for h in ss3["headers"]]
            rows3 = db.execute(
                f'SELECT * FROM "{ss3["table_name"]}" ORDER BY id').fetchall()
            tc = next((h for h in h3 if "合计" in h), h3[-1])
            vals3 = [dict(x)[tc] for x in rows3]
            check("显式启用：合计列重算回填", vals3 == [300, 700, 1100],
                  f"合计={vals3}（列={tc}）")
        finally:
            os.environ.pop("FINMOD_FORMULA_RECALC", None)

        # ── 4. 正常文件（有缓存值）不受影响 ──
        p2 = os.path.join(work, "正常.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S1"
        ws.append(["姓名", "金额"])
        ws.append(["张三", 100])
        ws.append(["李四", 200])
        wb.save(p2)
        wb.close()
        r2 = imp.import_workbook(proj["id"], "正常.xlsx", work)
        rows2 = db.execute(
            "SELECT * FROM data_%s_0 ORDER BY id" % r2["file_id"][:8]).fetchall()
        check("无公式文件不受影响", len(rows2) == 2,
              f"实际 {len(rows2)}")

    print()
    if FAIL:
        print("FAILED:", FAIL)
        return 1
    print("ALL PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
