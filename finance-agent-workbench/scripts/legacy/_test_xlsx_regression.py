"""Phase 17 P4 回归：.xlsx 写回链路不受 header_row 感知影响。

运行: python scripts/_test_xlsx_regression.py
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import storage.db as sdb  # noqa: E402
_tmp = tempfile.mkdtemp(prefix="xlsx_wb_reg_")
sdb.DB_PATH = os.path.join(_tmp, "test.db")

from app import create_app  # noqa: E402
from storage.project_store import ProjectStore  # noqa: E402
from storage.project_data_store import ProjectDataStore  # noqa: E402
from tools.project_data_importer import ProjectDataImporter  # noqa: E402
from tools.project_data_tools import export_project_data, update_rows  # noqa: E402

SRC = os.path.join(ROOT, "test-fixtures", "phase5", "待处理数据",
                   "销售提成汇总表2021-2026 0722.xlsx")
FAIL = []


def check(label: str, ok: bool, detail: str = "") -> None:
    print(("PASS" if ok else "FAIL"), label, detail)
    if not ok:
        FAIL.append(label)


def main() -> int:
    app = create_app()
    with app.app_context():
        from storage.db import get_db
        db = get_db()
        work = os.path.join(_tmp, "work")
        os.makedirs(work, exist_ok=True)
        dst = os.path.join(work, os.path.basename(SRC))
        shutil.copy2(SRC, dst)

        proj = ProjectStore(db).create("xlsx回归", work)
        imp = ProjectDataImporter(db)
        r = imp.import_workbook(proj["id"], os.path.basename(SRC), work)
        fid = r["file_id"]
        check("导入 .xlsx 成功", r.get("file_id") is not None,
              f"sheets={len(r['sheets'])}")

        store = ProjectDataStore(db)
        sheets = store.list_sheets(fid)
        sum_sheet = next(s for s in sheets if s["sheet_name"] == "汇总表")
        headers = [h["name"] for h in (sum_sheet.get("headers") or [])]
        print("  汇总表列名:", headers[:8], "... 共", len(headers))

        # 找一个可更新的行：第一列第一个非空值
        col0 = headers[0]
        from tools.project_data_tools import query_table
        q = query_table(file=os.path.basename(SRC), sheet="汇总表",
                        columns=[col0], limit=3, project_id=proj["id"],
                        work_dir=work)
        rows = q.get("rows") or q.get("data") or []
        key_val = rows[0][col0] if rows else "2021"
        target_col = headers[1]
        upd = update_rows(file=os.path.basename(SRC), sheet="汇总表",
                          filters={col0: key_val},
                          updates={target_col: 123456},
                          project_id=proj["id"], work_dir=work)
        check("update_rows 修改数据", upd.get("success"), str(upd.get("message")))

        exp = export_project_data(file=os.path.basename(SRC), mode="template",
                                  project_id=proj["id"], work_dir=work)
        check("template 写回成功", exp.get("success"), exp.get("message", str(exp)))
        check("写回格式识别为 xlsx", exp.get("format") == "xlsx", str(exp.get("format")))

        # 用 openpyxl 验证数据 + 表头位置
        import openpyxl
        wb = openpyxl.load_workbook(dst, data_only=True)
        ws = wb["汇总表"]
        found = False
        for row in ws.iter_rows(values_only=True):
            # 目标列是文本列（公司名），SQLite 会把数字存成 '123456' 字符串
            if row and row[0] == key_val and row[1] in (123456, "123456"):
                found = True
                break
        check("xlsx 写回数据验证", found, f"key={key_val}")
        check("xlsx 表头行保留", ws.cell(row=1, column=1).value is not None,
              f"row1col1={ws.cell(row=1, column=1).value}")
        wb.close()

        exp2 = export_project_data(file=os.path.basename(SRC), mode="precise",
                                   project_id=proj["id"], work_dir=work)
        check("precise 写回成功", exp2.get("success"), exp2.get("message", str(exp2)))

    print()
    if FAIL:
        print("FAILED:", FAIL)
        return 1
    print("ALL PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
