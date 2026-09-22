"""Excel → SQLite project data plane importer — Phase 14A.

Reads a workbook (.xlsx via openpyxl read_only+data_only, .xls via xlrd),
infers column types, materializes each sheet into a typed SQLite table
data_<file_id>_<idx>, and records metadata into project_data_files /
project_data_sheets via ProjectDataStore.

Design (PRD §3.12.4):
  - openpyxl read_only=True + data_only=True → formulas materialized as
    values (noted in meta_json).
  - SHA-256 file hash for idempotent re-import and later source-change
    detection (Phase 14D).
  - Same file + same hash → reuse existing record (idempotent).
"""

from __future__ import annotations

import hashlib
import json
import os
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable

from storage.project_data_store import ProjectDataStore
from tools.table_layout_detector import Layout, audit_and_correct, detect_layout

# Per-sheet row cap: oversized sheets are truncated to protect the DB.
MAX_ROWS_PER_SHEET = 200_000
# Sample values recorded per column in headers_json (2-3 per PRD §3.12.5).
SAMPLE_LIMIT = 3


class NeedAnchorError(Exception):
    """Raised when a sheet's header row could not be auto-detected with
    sufficient confidence and no user anchor was provided.

    Carries the probe result (per-sheet layouts + preview rows) so the
    caller can prompt the user to confirm the header row.
    """

    def __init__(self, file_hash: str, layouts: dict[str, Layout],
                 previews: dict[str, list[list]]) -> None:
        super().__init__("需要用户确认表头行")
        self.file_hash = file_hash
        self.layouts = layouts
        self.previews = previews


def _resolve_work_path(work_dir: str, rel_path: str) -> str:
    """Resolve a project-relative path inside the work directory."""
    work = Path(work_dir).resolve()
    full = (work / rel_path).resolve()
    if str(full) != str(work) and not str(full).startswith(str(work) + os.sep):
        raise ValueError(f"路径超出工作目录范围: {rel_path}")
    return str(full)


def sha256_file(path: str) -> str:
    """Compute SHA-256 of a file in 1 MiB chunks."""
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _sanitize_column(name: Any, used: set[str]) -> str:
    """Sanitize a raw header cell into a safe, unique SQLite column name."""
    raw = "" if name is None else str(name)
    base = re.sub(r"[^\w\u4e00-\u9fff]+", "_", raw).strip("_")
    if not base:
        base = "Col"
    candidate = base
    i = 1
    while candidate in used:
        i += 1
        candidate = f"{base}_{i}"
    used.add(candidate)
    return candidate


def _sqlite_type(value: Any) -> str:
    """Logical type of a single cell (int/float/date/text)."""
    if isinstance(value, bool):
        return "int"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, (datetime, date, time)):
        return "date"
    return "text"


def _promote_type(current: str, new: str) -> str:
    """Merge two per-column inferred types (widening)."""
    if current == new:
        return current
    if "text" in (current, new):
        return "text"
    if "date" in (current, new):
        return "date"
    if {current, new} == {"int", "float"}:
        return "float"
    return "text"


def _ddl_type(logical: str) -> str:
    """Map logical type to SQLite column DDL type."""
    if logical == "int":
        return "INTEGER"
    if logical == "float":
        return "REAL"
    if logical == "date":
        return "TEXT"
    return "TEXT"


def _norm_value(value: Any) -> Any:
    """Normalize a cell for storage (ISO datetimes, int floats, bool→int)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


class ProjectDataImporter:
    """Materialize one Excel workbook into the project data plane."""

    def __init__(self, db) -> None:
        self.db = db
        self.store = ProjectDataStore(db)

    # ------------------------------------------------------------------
    # Public entry
    # ------------------------------------------------------------------
    def read_workbook_rows(self, full: str, ext: str) -> dict[str, list[list]]:
        """Read every sheet into an in-memory list of rows (capped).

        Returns {sheet_name: [[cell, ...], ...]}.  Used both for probing
        (detect_layout) and for materialization with a user anchor.
        """
        if ext == ".xls":
            import xlrd  # lazy: .xls legacy support
            wb = xlrd.open_workbook(full)
            out: dict[str, list[list]] = {}
            for name in wb.sheet_names():
                sh = wb.sheet_by_name(name)
                out[name] = [
                    [self._xls_cell(sh, r, c) for c in range(sh.ncols)]
                    for r in range(min(sh.nrows, MAX_ROWS_PER_SHEET))
                ]
            return out

        import openpyxl  # lazy: bundled offline
        wb = openpyxl.load_workbook(full, read_only=True, data_only=True)
        out = {}
        try:
            for name in wb.sheetnames:
                ws = wb[name]
                rows = []
                for raw in ws.iter_rows(values_only=True):
                    if len(rows) >= MAX_ROWS_PER_SHEET:
                        break
                    rows.append(list(raw))
                # Phase 17 P11：公式兜底重算——第三方生成的无缓存值公式文件，
                # data_only 读到 None；用 formulas 库重算回填（零开销门：无 None/无公式则跳过）
                if rows:
                    try:
                        from tools.formula_fallback import formula_fallback_rows
                        rows = formula_fallback_rows(full, rows, name)
                    except Exception:  # noqa: BLE001
                        pass  # 兜底失败保持原样（数据平面仍可用）
                out[name] = rows
        finally:
            wb.close()
        return out

    def import_workbook(
        self,
        project_id: str,
        rel_path: str,
        work_dir: str,
        anchor_rows: dict[str, int] | None = None,
        always_anchor: bool = False,
        progress: Callable[[int, int, str], None] | None = None,
    ) -> dict:
        """Import a workbook.  Returns {"file_id", "reused", "file", "sheets"}.

        anchor_rows: {sheet_name: header_row} confirmed by the user.
        When a sheet's auto-detected confidence is low and no anchor is
        given for it, NeedAnchorError is raised with the probe result.

        always_anchor: True 时无论置信度高低都要求用户确认表头行
        （应用「每次上传都确认变量行/数据行」场景）。
        """
        full = _resolve_work_path(work_dir, rel_path)
        if not os.path.isfile(full):
            raise ValueError(f"文件不存在: {rel_path}")
        ext = os.path.splitext(full)[1].lower()
        if ext not in {".xlsx", ".xlsm", ".xls"}:
            raise ValueError(f"仅支持 Excel 文件: {rel_path}")

        file_hash = sha256_file(full)

        # Idempotent: same project + same hash (content unchanged) → reuse.
        existing = self.store.get_file_by_hash(project_id, file_hash)
        if existing:
            return {"file_id": existing["id"], "reused": True,
                    "file": existing, "sheets": self.store.list_sheets(existing["id"])}

        # Same source path already imported:
        #   - same hash → reuse (no duplicate rows)
        #   - different hash (Phase 14D: source file changed) → replace the
        #     stale record + data tables with a fresh import.
        dup = self.store.find_file_by_source(project_id, rel_path)
        if dup:
            if dup.get("file_hash") == file_hash:
                return {"file_id": dup["id"], "reused": True,
                        "file": dup, "sheets": self.store.list_sheets(dup["id"])}
            self.store.delete_file(dup["id"])

        # ── Read all rows once, probe layouts, apply anchors ──
        sheets_rows = self.read_workbook_rows(full, ext)
        layouts = {name: detect_layout(rows)
                   for name, rows in sheets_rows.items()}
        anchor_rows = anchor_rows or {}
        for name, hr in anchor_rows.items():
            if name == "__all__":
                # 通配锚点：应用到所有低置信 sheet（LLM 从用户处获取的
                # 通用表头行号，多数财务文件共用同一布局）。
                for nm, lay in list(layouts.items()):
                    if lay.is_confident or nm not in sheets_rows:
                        continue
                    rows = sheets_rows[nm]
                    lay.header_row = hr
                    lay.data_start_row = hr + 1
                    lay.confidence = 1.0
                    _used: set[str] = set()
                    lay.columns = [_sanitize_column(c, _used)
                                   for c in rows[hr]]
                continue
            if name in layouts and name in sheets_rows:
                rows = sheets_rows[name]
                lay = layouts[name]
                lay.header_row = hr
                lay.data_start_row = hr + 1
                lay.confidence = 1.0
                _used: set[str] = set()
                lay.columns = [_sanitize_column(c, _used)
                               for c in rows[hr]]

        # ── Col audit & self-correction (Phase 14E #2) ──
        # Blank header cells become "Col" columns.  Try to recover real
        # names via cross-sheet isomorphic fill + neighbor-derived naming.
        # Unrecovered Cols are reported so the caller can warn the user.
        uncorrected = audit_and_correct(sheets_rows, layouts)
        uncorrected = {
            name: idxs for name, idxs in uncorrected.items() if idxs
        }

        # ── 需要用户确认表头的 sheet ──
        # 默认：低置信度才问；always_anchor=True：全部 sheet 都问（每次上传确认）
        if always_anchor:
            low = {
                name: lay for name, lay in layouts.items()
                if name in sheets_rows and name not in anchor_rows
            }
        else:
            low = {
                name: lay for name, lay in layouts.items()
                if not lay.is_confident and name not in anchor_rows
            }
        if low:
            raise NeedAnchorError(
                file_hash=file_hash,
                layouts=layouts,
                previews={name: sheets_rows[name][:8] for name in low},
            )

        file_name = os.path.basename(full)
        file_rec = self.store.create_file(
            project_id, rel_path, file_name, file_hash, ext.lstrip(".")
        )
        file_id = file_rec["id"]
        try:
            sheets_meta: list[dict] = []
            total = max(len(sheets_rows), 1)
            for idx, (name, rows) in enumerate(sheets_rows.items()):
                meta = self._materialize_sheet(
                    file_id, idx, name, rows, layouts[name]
                )
                sheets_meta.append(meta)
                if progress:
                    progress(idx + 1, total, name)
        except Exception as exc:
            self.store.update_file(file_id, status="failed",
                                   error_message=str(exc)[:500])
            self.store.cleanup_sheets(file_id)
            raise

        total_rows = sum(s["row_count"] for s in sheets_meta)
        layout_meta = {
            name: {"header_row": lay.header_row,
                   "confidence": lay.confidence,
                   "anchored_by": "user" if name in anchor_rows else "auto"}
            for name, lay in layouts.items()
        }
        meta = {"file_size": os.path.getsize(full),
                "formulas_materialized": ext != ".xls",
                "layouts": layout_meta,
                "uncorrected_cols": uncorrected,
                "col_audit": "ok" if not uncorrected else "has_unnamed"}
        self.store.update_file(
            file_id,
            status="done",
            error_message=None,
            sheet_count=len(sheets_meta),
            total_rows=total_rows,
            meta_json=json.dumps(meta, ensure_ascii=False),
        )
        return {"file_id": file_id, "reused": False,
                "file": self.store.get_file(file_id), "sheets": sheets_meta}

    # ------------------------------------------------------------------
    # Workbook readers
    # ------------------------------------------------------------------
    def _xls_cell(self, sh, r: int, c: int) -> Any:
        import xlrd
        cell = sh.cell(r, c)
        v = cell.value
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(v, sh.book.datemode).isoformat()
            except Exception:
                return v
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return int(v)
        if isinstance(v, float) and v.is_integer():
            return int(v)
        return v

    # ------------------------------------------------------------------
    # Sheet materialization (anchor-aware)
    # ------------------------------------------------------------------
    def _materialize_sheet(
        self,
        file_id: str,
        sheet_index: int,
        sheet_name: str,
        rows: list[list],
        layout: Layout,
    ) -> dict:
        table_name = f"data_{file_id[:8]}_{sheet_index}"
        header_row = layout.header_row
        data_start = layout.data_start_row
        data_end = layout.data_end_row

        if header_row >= len(rows):
            raise ValueError(f"工作表 {sheet_name} 表头超出范围")

        used: set[str] = set()
        # Prefer the detector's (possibly self-corrected) column names; fall
        # back to re-sanitizing the raw header row for sheets materialized
        # with a user anchor (layout.columns may be empty there).
        if layout.columns:
            columns = [_sanitize_column(c, used) for c in layout.columns]
        else:
            columns = [_sanitize_column(c, used) for c in rows[header_row]]
        col_types: list[str | None] = [None] * len(columns)
        if not columns or all(name == "Col" for name in columns):
            raise ValueError(f"工作表 {sheet_name} 表头为空")

        # Data region: data_start .. data_end (data_end inclusive, -1 → end)
        end = data_end
        if end < 0 or end >= len(rows):
            end = len(rows) - 1

        data_rows: list[list] = []
        for idx in range(data_start, end + 1):
            if len(data_rows) >= MAX_ROWS_PER_SHEET:
                break
            raw = rows[idx]
            row = list(raw)[: len(columns)]
            while len(row) < len(columns):
                row.append(None)
            # Phase 17 P6：跳过数据区内的全空行（财务表常用空行做视觉分隔），
            # 避免产生全 null 数据行；不改变行号锚定（写回仍按行序）。
            if all(v is None or v == "" for v in row):
                continue
            for i, v in enumerate(row):
                t = _sqlite_type(v)
                col_types[i] = t if col_types[i] is None else _promote_type(col_types[i], t)
            data_rows.append([_norm_value(v) for v in row])

        # Columns with no non-null values fall back to text.
        col_types = [t or "text" for t in col_types]

        col_defs = [
            {"name": columns[i], "type": _ddl_type(col_types[i])}
            for i in range(len(columns))
        ]
        self.store.create_data_table(table_name, col_defs)
        self.store.insert_rows(table_name, columns, data_rows)

        headers = self._build_headers(columns, col_types, data_rows)
        self.store.create_sheet(
            file_id, sheet_name, sheet_index, table_name,
            len(data_rows), len(columns), headers,
        )
        return {
            "sheet_name": sheet_name,
            "sheet_index": sheet_index,
            "table_name": table_name,
            "row_count": len(data_rows),
            "col_count": len(columns),
        }

    def _build_headers(
        self, columns: list[str], col_types: list[str],
        data_rows: list[list],
    ) -> list[dict]:
        n = len(data_rows)
        headers: list[dict] = []
        for i, name in enumerate(columns):
            samples: list[str] = []
            non_null = 0
            for row in data_rows:
                v = row[i] if i < len(row) else None
                if v is not None and v != "":
                    non_null += 1
                    if len(samples) < SAMPLE_LIMIT:
                        samples.append(str(v))
            headers.append({
                "name": name,
                "type": col_types[i],
                "non_null_rate": round(non_null / n, 4) if n else 1.0,
                "samples": samples,
            })
        return headers
